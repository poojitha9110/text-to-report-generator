import streamlit as st
import ollama

from reportlab.lib.pagesizes import A4
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.enums import TA_CENTER

from io import BytesIO

from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

from pypdf import PdfReader


# ==========================================
# PAGE CONFIGURATION
# ==========================================

st.set_page_config(
    page_title="Text-to-Report Generator",
    page_icon="📄",
    layout="wide"
)


# ==========================================
# HEADER
# ==========================================

st.title("📄 Text-to-Report Generator")

st.markdown(
    "### Convert your text or documents into professional reports using AI."
)

st.info(
    "Upload a TXT, PDF, or DOCX file, or enter your text manually."
)


# ==========================================
# FILE UPLOAD
# ==========================================

uploaded_file = st.file_uploader(
    "📁 Upload a file",
    type=["txt", "pdf", "docx"]
)


# ==========================================
# MANUAL TEXT INPUT
# ==========================================

text = st.text_area(
    "✍️ Enter your text",
    height=250,
    placeholder="Enter or paste your text here..."
)


# ==========================================
# EXTRACT TEXT FROM FILE
# ==========================================

file_text = ""

if uploaded_file is not None:

    try:

        file_name = uploaded_file.name.lower()

        # ------------------------------
        # TXT FILE
        # ------------------------------

        if file_name.endswith(".txt"):

            file_text = uploaded_file.read().decode(
                "utf-8",
                errors="ignore"
            )


        # ------------------------------
        # PDF FILE
        # ------------------------------

        elif file_name.endswith(".pdf"):

            pdf_reader = PdfReader(uploaded_file)

            for page in pdf_reader.pages:

                extracted_text = page.extract_text()

                if extracted_text:

                    file_text += extracted_text + "\n"


        # ------------------------------
        # DOCX FILE
        # ------------------------------

        elif file_name.endswith(".docx"):

            document = Document(uploaded_file)

            for paragraph in document.paragraphs:

                if paragraph.text.strip():

                    file_text += paragraph.text + "\n"


        # ------------------------------
        # DISPLAY EXTRACTED TEXT
        # ------------------------------

        if file_text.strip():

            st.success(
                f"✅ File uploaded successfully: {uploaded_file.name}"
            )

            st.text_area(
                "📋 Extracted Text",
                file_text,
                height=200
            )

        else:

            st.warning(
                "⚠️ No readable text was found in the uploaded file."
            )


    except Exception as e:

        st.error(
            "⚠️ Error while reading the uploaded file."
        )

        st.write(e)


# ==========================================
# REPORT SETTINGS
# ==========================================

st.subheader("⚙️ Report Settings")


# ==========================================
# REPORT TYPE AND LENGTH
# ==========================================

col1, col2 = st.columns(2)


with col1:

    report_type = st.selectbox(
        "📑 Select Report Type",
        [
            "General Report",
            "Academic Report",
            "Project Report",
            "Technical Report",
            "Business Report"
        ]
    )


with col2:

    report_length = st.selectbox(
        "📏 Select Report Length",
        [
            "Short",
            "Medium",
            "Detailed"
        ]
    )


# ==========================================
# GENERATE REPORT BUTTON
# ==========================================

st.divider()


generate_button = st.button(
    "🚀 Generate Report",
    use_container_width=True
)


if generate_button:

    # ======================================
    # SELECT INPUT TEXT
    # ======================================

    if file_text.strip():

        input_text = file_text

    elif text.strip():

        input_text = text

    else:

        input_text = ""


    # ======================================
    # CHECK INPUT
    # ======================================

    if input_text.strip() == "":

        st.warning(
            "⚠️ Please enter text or upload a file."
        )


    else:

        try:

            # ==================================
            # OLLAMA REPORT GENERATION
            # ==================================

            with st.spinner(
                "🤖 Generating your report..."
            ):

                response = ollama.chat(

                    model="llama3.2",

                    messages=[

                        {

                            "role": "user",

                            "content": f"""
Create a professional {report_type}
with {report_length} length from the following text.

Create the report using exactly these sections:

TITLE
INTRODUCTION
MAIN CONTENT
CONCLUSION

Requirements:

- Use clear and simple English.
- Correct grammar and spelling.
- Organize the information properly.
- Do not add unrelated information.
- Do not use Markdown symbols such as # or *.
- Keep each section clearly separated.
- For Short: keep the report brief.
- For Medium: provide moderate detail.
- For Detailed: provide comprehensive detail.

Text:

{input_text}
"""

                        }

                    ]

                )


            # ==================================
            # GET GENERATED REPORT
            # ==================================

            report = response["message"]["content"]


            # ==================================
            # DISPLAY GENERATED REPORT
            # ==================================

            st.divider()

            st.subheader(
                "📑 Generated Report"
            )

            st.write(report)


            # ==================================
            # DOWNLOAD SECTION
            # ==================================

            st.subheader(
                "📥 Download Report"
            )


            # ==================================
            # TXT DOWNLOAD
            # ==================================

            st.download_button(

                label="📄 Download TXT",

                data=report,

                file_name="generated_report.txt",

                mime="text/plain",

                use_container_width=True

            )


            # ==================================
            # CREATE PDF
            # ==================================

            pdf_buffer = BytesIO()


            pdf_doc = SimpleDocTemplate(

                pdf_buffer,

                pagesize=A4,

                rightMargin=50,

                leftMargin=50,

                topMargin=50,

                bottomMargin=50

            )


            styles = getSampleStyleSheet()


            title_style = styles["Title"]

            title_style.alignment = TA_CENTER


            heading_style = styles["Heading2"]

            body_style = styles["BodyText"]


            story = []


            # PDF title

            story.append(

                Paragraph(

                    "Generated Report",

                    title_style

                )

            )


            story.append(

                Spacer(1, 20)

            )


            # Report type

            story.append(

                Paragraph(

                    f"Report Type: {report_type}",

                    heading_style

                )

            )


            story.append(

                Spacer(1, 15)

            )


            # Report content

            for line in report.split("\n"):

                line = line.strip()


                if not line:

                    continue


                upper_line = line.upper()


                if upper_line in [

                    "TITLE",

                    "INTRODUCTION",

                    "MAIN CONTENT",

                    "CONCLUSION"

                ]:

                    story.append(

                        Paragraph(

                            line,

                            heading_style

                        )

                    )


                else:

                    story.append(

                        Paragraph(

                            line,

                            body_style

                        )

                    )


                story.append(

                    Spacer(1, 8)

                )


            # Build PDF

            pdf_doc.build(story)


            pdf_data = pdf_buffer.getvalue()


            # ==================================
            # PDF DOWNLOAD
            # ==================================

            st.download_button(

                label="📕 Download PDF",

                data=pdf_data,

                file_name="generated_report.pdf",

                mime="application/pdf",

                use_container_width=True

            )


            # ==================================
            # CREATE DOCX
            # ==================================

            word_buffer = BytesIO()


            word_doc = Document()


            # DOCX title

            title = word_doc.add_heading(

                "Generated Report",

                level=0

            )


            title.alignment = (

                WD_ALIGN_PARAGRAPH.CENTER

            )


            # Report type

            type_paragraph = (

                word_doc.add_paragraph()

            )


            type_run = (

                type_paragraph.add_run(

                    f"Report Type: {report_type}"

                )

            )


            type_run.bold = True

            type_run.font.size = Pt(12)


            # DOCX report content

            for line in report.split("\n"):

                line = line.strip()


                if not line:

                    continue


                upper_line = line.upper()


                if upper_line in [

                    "TITLE",

                    "INTRODUCTION",

                    "MAIN CONTENT",

                    "CONCLUSION"

                ]:

                    word_doc.add_heading(

                        line,

                        level=1

                    )


                else:

                    paragraph = (

                        word_doc.add_paragraph()

                    )


                    run = (

                        paragraph.add_run(line)

                    )


                    run.font.size = Pt(11)


            # Save DOCX

            word_doc.save(word_buffer)


            word_data = word_buffer.getvalue()


            # ==================================
            # DOCX DOWNLOAD
            # ==================================

            st.download_button(

                label="📝 Download DOCX",

                data=word_data,

                file_name="generated_report.docx",

                mime=(
                    "application/vnd.openxmlformats-officedocument."
                    "wordprocessingml.document"
                ),

                use_container_width=True

            )


        # ======================================
        # OLLAMA ERROR HANDLING
        # ======================================

        except Exception as e:

            error_message = str(e).lower()


            if (

                "connect" in error_message

                or "connection" in error_message

                or "ollama" in error_message

            ):

                st.error(

                    "⚠️ Ollama is not running. "
                    "Please start Ollama and try again."

                )


            else:

                st.error(

                    "⚠️ Error while generating the report."

                )

                st.write(e)